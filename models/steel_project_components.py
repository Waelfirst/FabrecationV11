# -*- coding: utf-8 -*-
# models/steel_project_components.py - COMPLETE FILE
# Copy this entire file and replace your existing steel_project_components.py

from odoo import models, fields, api
import base64
import io
import logging
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class SteelProjectComponents(models.Model):
    _name = 'steel.project.components'
    _description = 'Steel Project Components'
    _rec_name = 'display_name'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    project_id = fields.Many2one(
        'steel.project',
        string='Project',
        required=True,
        ondelete='cascade'
    )
    product_id = fields.Many2one(
        'steel.project.product',
        string='Product',
        required=True,
        ondelete='cascade'
    )
    display_name = fields.Char(
        string='Name',
        compute='_compute_display_name',
        store=True
    )

    # Relations
    part_ids = fields.One2many(
        'steel.product.part',
        'component_id',
        string='Product Parts'
    )
    stage_ids = fields.One2many(
        'steel.production.stage',
        'component_id',
        string='Production Stages'
    )

    # Computed counts
    total_parts = fields.Integer(
        string='Total Parts',
        compute='_compute_totals',
        store=True
    )
    total_stages = fields.Integer(
        string='Total Stages',
        compute='_compute_totals',
        store=True
    )
    total_parts_quantity = fields.Float(
        string='Total Parts Quantity',
        compute='_compute_totals',
        store=True
    )

    # Total costs
    total_parts_cost = fields.Monetary(
        string='Total Parts Cost',
        compute='_compute_cost_totals',
        store=True,
        currency_field='currency_id'
    )
    total_stages_cost = fields.Monetary(
        string='Total Stages Cost',
        compute='_compute_cost_totals',
        store=True,
        currency_field='currency_id'
    )
    total_component_cost = fields.Monetary(
        string='Total Component Cost',
        compute='_compute_cost_totals',
        store=True,
        currency_field='currency_id'
    )
    currency_id = fields.Many2one(
        'res.currency',
        string='Currency',
        default=lambda self: self.env.company.currency_id
    )

    # Import fields for parts
    import_file = fields.Binary(string='Import Excel File', attachment=False)
    import_filename = fields.Char(string='Filename')

    # Import fields for stages
    import_stage_file = fields.Binary(string='Import Stage Excel File', attachment=False)
    import_stage_filename = fields.Char(string='Stage Filename')

    # Auto-create toggle
    auto_create_missing_records = fields.Boolean(
        string='Auto-Create Missing Records',
        default=True,
        help='Automatically create materials, workcenters, operations, workers, and machines if not found during import'
    )

    # Status tracking
    state = fields.Selection([
        ('draft', 'Draft'),
        ('confirmed', 'Confirmed'),
        ('in_production', 'In Production'),
        ('done', 'Done'),
    ], string='Status', default='draft', tracking=True)

    @api.depends('project_id', 'project_id.name', 'product_id', 'product_id.product_name')
    def _compute_display_name(self):
        for record in self:
            if record.project_id and record.product_id:
                product_name = record.product_id.product_name or 'Unnamed Product'
                record.display_name = f"{record.project_id.name} - {product_name}"
            else:
                record.display_name = "New Component"

    @api.depends('part_ids', 'stage_ids', 'part_ids.total_quantity')
    def _compute_totals(self):
        for record in self:
            record.total_parts = len(record.part_ids)
            record.total_stages = len(record.stage_ids)
            record.total_parts_quantity = sum(record.part_ids.mapped('total_quantity'))

    @api.depends('part_ids.total_cost', 'stage_ids.estimated_labor_cost')
    def _compute_cost_totals(self):
        for record in self:
            record.total_parts_cost = sum(record.part_ids.mapped('total_cost'))
            record.total_stages_cost = sum(record.stage_ids.mapped('estimated_labor_cost'))
            record.total_component_cost = record.total_parts_cost + record.total_stages_cost

    # ========================================
    # AUTO-CREATE HELPER METHODS
    # ========================================

    def _create_missing_material(self, material_name):
        """Create a new material/product if not found"""
        try:
            material = self.env['product.product'].create({
                'name': material_name,
                'type': 'product',
                'categ_id': self.env.ref('product.product_category_all').id,
                'uom_id': self.env.ref('uom.product_uom_unit').id,
                'uom_po_id': self.env.ref('uom.product_uom_unit').id,
                'standard_price': 0.0,
                'list_price': 0.0,
            })
            _logger.info(f'✅ Auto-created material: {material_name}')
            return material
        except Exception as e:
            _logger.error(f'❌ Failed to create material {material_name}: {str(e)}')
            return False

    def _create_missing_workcenter(self, workcenter_name):
        """Create a new workcenter if not found"""
        try:
            workcenter = self.env['mrp.workcenter'].create({
                'name': workcenter_name,
                'code': workcenter_name[:10].upper().replace(' ', '_'),
                'time_efficiency': 100.0,
                'costs_hour': 50.0,
            })
            _logger.info(f'✅ Auto-created workcenter: {workcenter_name}')
            return workcenter
        except Exception as e:
            _logger.error(f'❌ Failed to create workcenter {workcenter_name}: {str(e)}')
            return False

    def _create_missing_operation(self, operation_name, workcenter_id=False):
        """Create a new manufacturing operation if not found"""
        try:
            # Find or create a default routing
            routing = self.env['mrp.routing'].search([], limit=1)
            if not routing:
                routing = self.env['mrp.routing'].create({
                    'name': 'Default Routing',
                })

            operation = self.env['mrp.routing.workcenter'].create({
                'name': operation_name,
                'routing_id': routing.id,
                'workcenter_id': workcenter_id if workcenter_id else False,
                'time_cycle_manual': 1.0,
                'sequence': 10,
            })
            _logger.info(f'✅ Auto-created operation: {operation_name}')
            return operation
        except Exception as e:
            _logger.error(f'❌ Failed to create operation {operation_name}: {str(e)}')
            return False

    def _create_missing_worker(self, worker_name):
        """Create a new employee/worker if not found"""
        try:
            worker = self.env['hr.employee'].create({
                'name': worker_name,
                'job_title': 'Worker',
            })
            _logger.info(f'✅ Auto-created worker: {worker_name}')
            return worker
        except Exception as e:
            _logger.error(f'❌ Failed to create worker {worker_name}: {str(e)}')
            return False

    def _create_missing_machine(self, machine_name):
        """Create a new machine/equipment if not found"""
        try:
            machine = self.env['maintenance.equipment'].create({
                'name': machine_name,
                'equipment_assign_to': 'other',
            })
            _logger.info(f'✅ Auto-created machine: {machine_name}')
            return machine
        except Exception as e:
            _logger.error(f'❌ Failed to create machine {machine_name}: {str(e)}')
            return False

    # ========================================
    # PARTS IMPORT WITH AUTO-CREATE
    # ========================================

    def action_import_parts_from_excel(self):
        """Import parts from Excel with auto-creation"""
        self.ensure_one()

        if not self.import_file:
            raise UserError("Please upload an Excel file first.")

        try:
            file_data = base64.b64decode(self.import_file)
            imported_count = 0
            created_materials_count = 0
            errors = []
            warnings = []

            try:
                from openpyxl import load_workbook
                workbook = load_workbook(filename=io.BytesIO(file_data), read_only=True, data_only=True)
                sheet = workbook.active

                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                    try:
                        part_name = str(row[0]).strip() if row[0] else ''
                        part_number = str(row[1]).strip() if len(row) > 1 and row[1] else ''
                        material_name = str(row[2]).strip() if len(row) > 2 and row[2] else ''
                        quantity = float(row[3]) if len(row) > 3 and row[3] else 1.0
                        notes = str(row[4]).strip() if len(row) > 4 and row[4] else ''

                        if not part_name or not material_name:
                            continue

                        # Search for material
                        material = self.env['product.product'].search([
                            ('name', 'ilike', material_name)
                        ], limit=1)

                        # Auto-create material if not found
                        if not material:
                            if self.auto_create_missing_records:
                                material = self._create_missing_material(material_name)
                                if material:
                                    created_materials_count += 1
                                    warnings.append(f"Row {row_idx}: ✅ Created new material '{material_name}'")
                                else:
                                    errors.append(f"Row {row_idx}: ❌ Failed to create material '{material_name}'")
                                    continue
                            else:
                                errors.append(f"Row {row_idx}: ❌ Material '{material_name}' not found")
                                continue

                        # Create part
                        self.env['steel.product.part'].create({
                            'component_id': self.id,
                            'part_name': str(part_name),
                            'part_number': str(part_number) if part_number else False,
                            'material_id': material.id,
                            'quantity': quantity,
                            'notes': str(notes) if notes else False,
                        })
                        imported_count += 1

                    except Exception as e:
                        errors.append(f"Row {row_idx}: {str(e)}")
                        continue

            except ImportError:
                raise UserError("Please install 'openpyxl' library to import .xlsx files.\nRun: pip install openpyxl")

            # Build result message
            message = f"✅ Successfully imported {imported_count} parts."

            if created_materials_count > 0:
                message += f"\n\n🆕 Auto-created {created_materials_count} new material(s)."

            if warnings:
                message += f"\n\n⚠️ Auto-Created Items:\n" + "\n".join(warnings[:15])
                if len(warnings) > 15:
                    message += f"\n... and {len(warnings) - 15} more"

            if errors:
                message += f"\n\n❌ Errors:\n" + "\n".join(errors[:10])
                if len(errors) > 10:
                    message += f"\n... and {len(errors) - 10} more errors"

            # Clear import fields
            self.write({
                'import_file': False,
                'import_filename': False
            })

            return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {
                    'title': 'Parts Import Complete',
                    'message': message,
                    'type': 'success' if imported_count > 0 else 'warning',
                    'sticky': True,
                }
            }

        except Exception as e:
            raise UserError(f"Error reading Excel file: {str(e)}\n\nPlease make sure the file format is correct.")

    # ========================================
    # STAGES IMPORT WITH AUTO-CREATE AND MULTI-SELECT
    # ========================================

    def action_import_stages_from_excel(self):
        """
        Enhanced import stages from Excel with:
        - Multiple workers support (comma-separated)
        - Multiple machines support (comma-separated)
        - Duplicate operation prevention
        - Auto-creation of missing records
        """
        self.ensure_one()

        if not self.import_stage_file:
            raise UserError("Please upload an Excel file first.")

        try:
            file_data = base64.b64decode(self.import_stage_file)
            imported_count = 0
            skipped_duplicates = 0
            created_workcenters_count = 0
            created_operations_count = 0
            created_workers_count = 0
            created_machines_count = 0
            errors = []
            warnings = []

            try:
                from openpyxl import load_workbook
                workbook = load_workbook(filename=io.BytesIO(file_data), read_only=True, data_only=True)
                sheet = workbook.active

                _logger.info(f'📊 Starting Excel import - Sheet: {sheet.title}, Rows: {sheet.max_row}')

                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                    try:
                        # Skip empty rows
                        if not any(row):
                            continue

                        # Extract values from Excel
                        part_name = str(row[0]).strip() if row[0] else ''
                        workcenter_names = str(row[1]).strip() if len(row) > 1 and row[1] else ''
                        operation_names = str(row[2]).strip() if len(row) > 2 and row[2] else ''
                        operation_times = str(row[3]).strip() if len(row) > 3 and row[3] else ''
                        worker_names = str(row[4]).strip() if len(row) > 4 and row[4] else ''
                        machine_names = str(row[5]).strip() if len(row) > 5 and row[5] else ''
                        required_time = float(row[6]) if len(row) > 6 and row[6] else 0.0
                        hourly_rate = float(row[7]) if len(row) > 7 and row[7] else 50.0
                        notes = str(row[8]).strip() if len(row) > 8 and row[8] else ''

                        # Validate required fields
                        if not part_name:
                            warnings.append(f"Row {row_idx}: ⏭️ Skipped - No part name")
                            continue

                        if not operation_names:
                            errors.append(f"Row {row_idx}: ❌ Part '{part_name}' - No operations specified")
                            continue

                        _logger.info(f'Processing Row {row_idx}: Part={part_name}, Operations={operation_names}')

                        # Find the part
                        part = self.env['steel.product.part'].search([
                            ('component_id', '=', self.id),
                            ('part_name', 'ilike', part_name)
                        ], limit=1)

                        if not part:
                            errors.append(f"Row {row_idx}: ❌ Part '{part_name}' not found in component")
                            continue

                        # Calculate time from operation times
                        calculated_time = 0.0
                        if operation_times:
                            try:
                                times_list = str(operation_times).replace(' ', '').split(',')
                                for time_str in times_list:
                                    if time_str:
                                        calculated_time += float(time_str)
                            except ValueError:
                                warnings.append(f"Row {row_idx}: ⚠️ Invalid operation times format")

                        final_time = calculated_time if calculated_time > 0 else (
                            required_time if required_time > 0 else 1.0)

                        # Process each operation
                        operation_list = [op.strip() for op in operation_names.split(',') if op.strip()]

                        for operation_name in operation_list:
                            # CHECK FOR DUPLICATES
                            operation = self.env['mrp.routing.workcenter'].search([
                                ('name', 'ilike', operation_name)
                            ], limit=1)

                            # If operation exists, check for duplicates
                            if operation:
                                existing_stage = self.env['steel.production.stage'].search([
                                    ('component_id', '=', self.id),
                                    ('part_id', '=', part.id),
                                    ('operation_id', '=', operation.id)
                                ], limit=1)

                                if existing_stage:
                                    skipped_duplicates += 1
                                    warnings.append(
                                        f"Row {row_idx}: ⏭️ DUPLICATE SKIPPED - "
                                        f"Operation '{operation_name}' already exists for part '{part_name}'"
                                    )
                                    continue

                            # Create operation if not found
                            if not operation and self.auto_create_missing_records:
                                routing = self.env['mrp.routing'].search([], limit=1)
                                if not routing:
                                    routing = self.env['mrp.routing'].create({'name': 'Default Routing'})

                                # Find or create workcenter
                                workcenter_id = False
                                if workcenter_names:
                                    first_wc = workcenter_names.split(',')[0].strip()
                                    workcenter = self.env['mrp.workcenter'].search([
                                        ('name', 'ilike', first_wc)
                                    ], limit=1)

                                    if not workcenter:
                                        workcenter = self._create_missing_workcenter(first_wc)
                                        if workcenter:
                                            created_workcenters_count += 1
                                            warnings.append(f"Row {row_idx}: 🏭 Created workcenter '{first_wc}'")

                                    if workcenter:
                                        workcenter_id = workcenter.id

                                operation = self.env['mrp.routing.workcenter'].create({
                                    'name': operation_name,
                                    'routing_id': routing.id,
                                    'workcenter_id': workcenter_id,
                                    'time_cycle_manual': final_time,
                                    'sequence': 10,
                                })
                                created_operations_count += 1
                                warnings.append(f"Row {row_idx}: ⚙️ Created operation '{operation_name}'")

                            if not operation:
                                errors.append(f"Row {row_idx}: ❌ Operation '{operation_name}' not found")
                                continue

                            # PROCESS MULTIPLE WORKERS
                            worker_ids = []
                            if worker_names:
                                for worker_name in worker_names.split(','):
                                    worker_name = worker_name.strip()
                                    if worker_name:
                                        worker = self.env['hr.employee'].search([
                                            ('name', 'ilike', worker_name)
                                        ], limit=1)

                                        if not worker and self.auto_create_missing_records:
                                            worker = self._create_missing_worker(worker_name)
                                            if worker:
                                                created_workers_count += 1
                                                warnings.append(f"Row {row_idx}: 👷 Created worker '{worker_name}'")

                                        if worker:
                                            worker_ids.append(worker.id)
                                        else:
                                            warnings.append(f"Row {row_idx}: ⚠️ Worker '{worker_name}' not found")

                            # PROCESS MULTIPLE MACHINES
                            machine_ids = []
                            if machine_names:
                                for machine_name in machine_names.split(','):
                                    machine_name = machine_name.strip()
                                    if machine_name:
                                        machine = self.env['maintenance.equipment'].search([
                                            ('name', 'ilike', machine_name)
                                        ], limit=1)

                                        if not machine and self.auto_create_missing_records:
                                            machine = self._create_missing_machine(machine_name)
                                            if machine:
                                                created_machines_count += 1
                                                warnings.append(f"Row {row_idx}: 🔧 Created machine '{machine_name}'")

                                        if machine:
                                            machine_ids.append(machine.id)
                                        else:
                                            warnings.append(f"Row {row_idx}: ⚠️ Machine '{machine_name}' not found")

                            # CREATE STAGE WITH MULTI-SELECT
                            stage_vals = {
                                'component_id': self.id,
                                'part_id': part.id,
                                'operation_id': operation.id,
                                'required_time': final_time,
                                'hourly_rate': hourly_rate,
                                'description': notes if notes else f'Stage for {part.part_name} - {operation_name}',
                            }

                            # Add workers (many2many)
                            if worker_ids:
                                stage_vals['worker_ids'] = [(6, 0, worker_ids)]

                            # Add machines (many2many)
                            if machine_ids:
                                stage_vals['machine_ids'] = [(6, 0, machine_ids)]

                            # Create the stage
                            self.env['steel.production.stage'].create(stage_vals)
                            imported_count += 1

                            worker_info = f" with {len(worker_ids)} worker(s)" if worker_ids else ""
                            machine_info = f" and {len(machine_ids)} machine(s)" if machine_ids else ""
                            _logger.info(f"✅ Created: {part.part_name} - {operation_name}{worker_info}{machine_info}")

                    except Exception as e:
                        error_msg = f"Row {row_idx}: ❌ {str(e)}"
                        _logger.error(error_msg)
                        errors.append(error_msg)
                        continue

            except ImportError:
                raise UserError(
                    "Please install 'openpyxl' library to import .xlsx files.\n"
                    "Run: pip install openpyxl"
                )

            # BUILD RESULT MESSAGE
            message = f"✅ Successfully imported {imported_count} stage(s)."

            if skipped_duplicates > 0:
                message += f"\n⏭️ Skipped {skipped_duplicates} duplicate(s) (same operation + part combination)."

            if created_workcenters_count > 0:
                message += f"\n🏭 Auto-created {created_workcenters_count} workcenter(s)."
            if created_operations_count > 0:
                message += f"\n⚙️ Auto-created {created_operations_count} operation(s)."
            if created_workers_count > 0:
                message += f"\n👷 Auto-created {created_workers_count} worker(s)."
            if created_machines_count > 0:
                message += f"\n🔧 Auto-created {created_machines_count} machine(s)."

            if warnings:
                message += f"\n\n⚠️ Warnings ({len(warnings)}):\n" + "\n".join(warnings[:20])
                if len(warnings) > 20:
                    message += f"\n... and {len(warnings) - 20} more warnings"

            if errors:
                message += f"\n\n❌ Errors ({len(errors)}):\n" + "\n".join(errors[:10])
                if len(errors) > 10:
                    message += f"\n... and {len(errors) - 10} more errors"

            # Clear import fields
            self.write({
                'import_stage_file': False,
                'import_stage_filename': False
            })

            return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {
                    'title': 'Stage Import Complete',
                    'message': message,
                    'type': 'success' if imported_count > 0 else 'warning',
                    'sticky': True,
                }
            }

        except Exception as e:
            _logger.error(f"Excel import error: {str(e)}")
            raise UserError(f"Error reading Excel file: {str(e)}\n\nPlease ensure the file format is correct.")

    # ========================================
    # DOWNLOAD TEMPLATE METHODS
    # ========================================

    def action_download_parts_template(self):
        """Download Excel template for parts import"""
        self.ensure_one()

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment

            wb = Workbook()
            ws = wb.active
            ws.title = "Parts Import Template"

            headers = ['Part Name*', 'Part Number', 'Material Name*', 'Quantity per Unit*', 'Notes']
            ws.append(headers)

            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')

            ws.append(['Primary Column', 'COL-A1', 'Steel Column 300x300', 4, 'Main structural columns'])
            ws.append(['Main Beam Horizontal', 'BEAM-H1', 'Steel I-Beam 200mm', 8, 'Horizontal support beams'])
            ws.append(['Connection Plate', 'PLATE-001', 'Steel Plate 10mm', 50, 'Welding connection plates'])

            ws.column_dimensions['A'].width = 25
            ws.column_dimensions['B'].width = 15
            ws.column_dimensions['C'].width = 25
            ws.column_dimensions['D'].width = 20
            ws.column_dimensions['E'].width = 30

            ws.append([])
            ws.append(['INSTRUCTIONS:'])
            ws.append(['1. Fill in the required fields marked with * (asterisk)'])
            ws.append(['2. Part Name: Name of the part (e.g., "Primary Column")'])
            ws.append(['3. Part Number: Optional reference code (e.g., "COL-A1")'])
            ws.append(['4. Material Name: Material name - will be created automatically if not found!'])
            ws.append(['5. Quantity per Unit: How many parts needed per one product unit'])
            ws.append(['6. Notes: Optional additional information'])
            ws.append(['7. Delete the example rows before importing'])
            ws.append(['8. Save file and upload in the component form'])

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            file_data = base64.b64encode(output.read())

            attachment = self.env['ir.attachment'].create({
                'name': 'Parts_Import_Template.xlsx',
                'type': 'binary',
                'datas': file_data,
                'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            })

            return {
                'type': 'ir.actions.act_url',
                'url': f'/web/content/{attachment.id}?download=true',
                'target': 'new',
            }

        except ImportError:
            raise UserError("Please install 'openpyxl' library to generate Excel templates.\nRun: pip install openpyxl")

    def action_download_stages_template(self):
        """Download enhanced Excel template with multi-select support"""
        self.ensure_one()

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "Stages Import"

            # Define headers (with Machines column)
            headers = [
                'Part Name*',
                'Workcenters',
                'Operations*',
                'Operation Times (hrs)',
                'Workers (Multi)',
                'Machines (Multi)',
                'Required Time (hrs)',
                'Hourly Rate',
                'Notes'
            ]

            ws.append(headers)

            # Style header row
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = thin_border

            # Add example data
            if self.part_ids:
                for idx, part in enumerate(self.part_ids[:3]):
                    ws.append([
                        part.part_name,
                        'Cutting Station, Welding Station',
                        'Cut-1, Weld-1, Finish-1',
                        '2, 3, 1.5',
                        'John Welder, Mary Operator, Bob Finisher',
                        'CNC Cutter, Welding Machine A, Grinding Machine',
                        '',  # Auto-calculate
                        75.0,
                        f'Multiple stages for {part.part_name}'
                    ])
            else:
                # Generic examples
                ws.append([
                    'Primary Column',
                    'Cutting Station',
                    'Cut-1',
                    '2',
                    'John Operator, Mary Assistant',
                    'CNC Plasma Cutter, Material Handler',
                    '',
                    50.0,
                    'Cutting with 2 workers and 2 machines'
                ])
                ws.append([
                    'Primary Column',
                    'Welding Station',
                    'Weld-1',
                    '4',
                    'Bob Welder, Alice Helper, Tom Inspector',
                    'Welding Machine A, Welding Machine B, Grinder',
                    '',
                    75.0,
                    'Welding with 3 workers and 3 machines'
                ])
                ws.append([
                    'Main Beam',
                    'Assembly Station',
                    'Assemble-1',
                    '3',
                    'Charlie Assembler, Diana Finisher',
                    'Crane, Press Brake, Drill Press',
                    '',
                    60.0,
                    'Assembly with multiple workers and equipment'
                ])

            # Set column widths
            column_widths = [25, 35, 35, 25, 40, 40, 18, 15, 40]
            for idx, width in enumerate(column_widths, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = width

            # Add comprehensive instructions
            instruction_row = ws.max_row + 2
            ws.append([])
            ws.append(['📋 PRODUCTION STAGE IMPORT - COMPLETE GUIDE'])

            instructions = [
                '',
                '✅ REQUIRED FIELDS:',
                '• Part Name*: Must match an existing part',
                '• Operations*: MRP operation names',
                '',
                '👥 MULTIPLE WORKERS & MACHINES:',
                '• Workers: Comma-separated names (John, Mary, Bob)',
                '• Machines: Comma-separated names (CNC-1, Welder-A)',
                '• ALL listed workers/machines are assigned',
                '',
                '🚫 DUPLICATE PREVENTION:',
                '• Each operation can only be added ONCE per part',
                '• Duplicates are automatically skipped',
                '',
                '📊 AUTO-CALCULATED:',
                '• Required Time: Leave blank to auto-calculate',
                '',
                '🆕 AUTO-CREATE (when enabled):',
                '• Workcenters, Operations, Workers, Machines',
                '',
                f'📊 EXISTING PARTS ({len(self.part_ids)}):',
            ]

            if self.part_ids:
                instructions.append('')
                for part in self.part_ids:
                    material_info = f' ({part.material_id.name})' if part.material_id else ''
                    instructions.append(f'  ✅ {part.part_name}{material_info}')
            else:
                instructions.append('  ⚠️ No parts defined yet!')

            instructions.extend([
                '',
                '⚠️ IMPORTANT:',
                '• Delete example rows before importing',
                '• Save as .xlsx format',
                '• Upload in component form',
                '• Enable Auto-Create toggle if needed',
            ])

            for instruction in instructions:
                ws.append([instruction])

            # Style instructions
            instruction_font = Font(size=10)
            section_font = Font(bold=True, size=11, color="E67E22")

            for row in ws.iter_rows(min_row=instruction_row, max_row=ws.max_row):
                for cell in row:
                    if cell.value:
                        if cell.value.startswith(('📋', '✅', '👥', '🚫', '📊', '🆕', '⚠️')):
                            cell.font = section_font
                        else:
                            cell.font = instruction_font

            # Freeze panes
            ws.freeze_panes = 'A2'

            # Save
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            file_data = base64.b64encode(output.read())

            timestamp = fields.Datetime.now().strftime('%Y%m%d_%H%M%S')
            attachment = self.env['ir.attachment'].create({
                'name': f'Production_Stages_Template_{self.display_name}_{timestamp}.xlsx',
                'type': 'binary',
                'datas': file_data,
                'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            })

            return {
                'type': 'ir.actions.act_url',
                'url': f'/web/content/{attachment.id}?download=true',
                'target': 'new',
            }

        except ImportError:
            raise UserError("Please install 'openpyxl' library.\nRun: pip install openpyxl")

    # ========================================
    # OTHER ACTIONS
    # ========================================

    def action_create_stages_for_all_parts(self):
        """Create production stages for all parts in this component"""
        self.ensure_one()

        if not self.part_ids:
            raise UserError("No parts found! Please add parts first before creating stages.")

        existing_stages_count = len(self.stage_ids)
        default_workcenter = self.env['mrp.workcenter'].search([], limit=1)
        default_required_time = 1.0
        default_hourly_rate = 50.0

        created_count = 0
        skipped_count = 0

        for part in self.part_ids:
            try:
                existing_stage = self.stage_ids.filtered(lambda s: s.part_id.id == part.id)

                if existing_stage:
                    skipped_count += 1
                    continue

                stage_vals = {
                    'component_id': self.id,
                    'part_id': part.id,
                    'sequence': (existing_stages_count + created_count + 1) * 10,
                    'required_time': default_required_time,
                    'hourly_rate': default_hourly_rate,
                    'description': f'Production stage for {part.part_name}',
                }

                if default_workcenter:
                    stage_vals['workcenter_ids'] = [(6, 0, [default_workcenter.id])]

                self.env['steel.production.stage'].create(stage_vals)
                created_count += 1

            except Exception as e:
                _logger.warning(f'Failed to create stage for part {part.part_name}: {str(e)}')
                continue

        message = f"✅ Successfully created {created_count} production stage(s)."
        if skipped_count > 0:
            message += f"\n⚠️ Skipped {skipped_count} part(s) that already have stages."

        if created_count > 0:
            message += "\n\n💡 Tip: You can now customize the workcenter, time, and hourly rate for each stage."

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Stages Created Successfully' if created_count > 0 else 'No Stages Created',
                'message': message,
                'type': 'success' if created_count > 0 else 'info',
                'sticky': True,
            }
        }

    def action_clear_all_stages(self):
        """Delete all production stages for this component"""
        self.ensure_one()

        if not self.stage_ids:
            raise UserError("No stages to delete!")

        stage_count = len(self.stage_ids)

        try:
            self.stage_ids.unlink()

            return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {
                    'title': 'Stages Cleared',
                    'message': f'Successfully deleted {stage_count} stage(s).',
                    'type': 'info',
                    'sticky': False,
                }
            }
        except Exception as e:
            _logger.error(f'Failed to delete stages: {str(e)}')
            raise UserError(f'Failed to delete stages: {str(e)}')

    def action_calculate_all_stages_time(self):
        """Calculate time for all stages from their operations/workcenters"""
        self.ensure_one()

        if not self.stage_ids:
            raise UserError("No stages found! Please add stages first.")

        updated_count = 0
        skipped_count = 0
        total_time = 0.0

        for stage in self.stage_ids:
            stage_time = 0.0

            if stage.operation_id and stage.operation_id.time_cycle_manual:
                stage_time = stage.operation_id.time_cycle_manual

            if stage_time > 0:
                stage.required_time = stage_time
                total_time += stage_time
                updated_count += 1
            else:
                skipped_count += 1

        message = f"✅ Updated {updated_count} stage(s) with calculated time.\n"
        message += f"📊 Total project time: {total_time} hours\n"

        if skipped_count > 0:
            message += f"\n⚠️ Skipped {skipped_count} stage(s) with no time data."

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Bulk Time Calculation Complete',
                'message': message,
                'type': 'success',
                'sticky': True,
            }
        }

    def action_confirm(self):
        """Confirm component"""
        self.write({'state': 'confirmed'})

    def action_start_production(self):
        """Start production"""
        self.write({'state': 'in_production'})

    def action_done(self):
        """Mark as done"""
        self.write({'state': 'done'})

    def action_draft(self):
        """Reset to draft"""
        self.write({'state': 'draft'})

    def action_view_parts(self):
        """View parts of this component"""
        self.ensure_one()
        return {
            'name': 'Parts',
            'type': 'ir.actions.act_window',
            'res_model': 'steel.product.part',
            'view_mode': 'tree,form',
            'domain': [('component_id', '=', self.id)],
            'context': {'default_component_id': self.id},
        }

    def action_view_stages(self):
        """View stages of this component"""
        self.ensure_one()
        return {
            'name': 'Production Stages',
            'type': 'ir.actions.act_window',
            'res_model': 'steel.production.stage',
            'view_mode': 'tree,form',
            'domain': [('component_id', '=', self.id)],
            'context': {'default_component_id': self.id},
        }